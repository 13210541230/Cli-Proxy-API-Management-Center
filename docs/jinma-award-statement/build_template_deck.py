from pathlib import Path
import shutil
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parent
TEMPLATE = Path(r"D:/Downloads/2026年金码奖申报陈述PPT模板-XX部门-姓名.pptx")
OUTPUT = ROOT / "jinma-award-statement-2026-template.pptx"
BACKUP = ROOT / "jinma-award-statement-25page-archive.pptx"

RED = RGBColor(0xED, 0x00, 0x16)
BLUE = RGBColor(0x06, 0x67, 0xB1)
INK = RGBColor(0x22, 0x22, 0x22)
MUTED = RGBColor(0x66, 0x66, 0x66)
PALE = RGBColor(0xF5, 0xF7, 0xF9)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
DARK = RGBColor(0x18, 0x23, 0x2D)


def set_run(run, size=18, bold=False, color=INK, font="Microsoft YaHei", italic=False):
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color


def clear_frame(shape):
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(0.08)
    tf.margin_right = Inches(0.08)
    tf.margin_top = Inches(0.04)
    tf.margin_bottom = Inches(0.04)
    return tf


def set_plain(shape, text, size=18, bold=False, color=INK, align=None, font="Microsoft YaHei"):
    tf = clear_frame(shape)
    p = tf.paragraphs[0]
    p.alignment = align if align is not None else PP_ALIGN.LEFT
    r = p.add_run()
    r.text = text
    set_run(r, size=size, bold=bold, color=color, font=font)
    return tf


def replace_all_runs(shape, replacements):
    if not hasattr(shape, "text_frame"):
        return
    for p in shape.text_frame.paragraphs:
        for r in p.runs:
            for old, new in replacements.items():
                if old in r.text:
                    r.text = r.text.replace(old, new)


def replace_paragraph_placeholders(shape, values):
    if not hasattr(shape, "text_frame"):
        return
    value_index = 0
    for p in shape.text_frame.paragraphs:
        for r in p.runs:
            if "XXX" in r.text and value_index < len(values):
                r.text = r.text.replace("XXX", values[value_index], 1)
                value_index += 1


def replace_paragraphs(shape, values, size=20):
    if not hasattr(shape, "text_frame"):
        return
    paragraphs = shape.text_frame.paragraphs
    for index, value in enumerate(values):
        if index >= len(paragraphs):
            break
        paragraph = paragraphs[index]
        paragraph.clear()
        run = paragraph.add_run()
        run.text = value
        set_run(run, size=size, color=INK)


def add_rich_paragraph(tf, label, body, size=15, color=INK, after=4):
    p = tf.add_paragraph()
    p.alignment = PP_ALIGN.LEFT
    p.space_after = Pt(after)
    r1 = p.add_run(); r1.text = label
    set_run(r1, size=size, bold=True, color=color)
    r2 = p.add_run(); r2.text = body
    set_run(r2, size=size, color=color)
    return p


def add_text(slide, x, y, w, h, text, size=16, bold=False, color=INK, align=PP_ALIGN.LEFT, font="Microsoft YaHei"):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = clear_frame(shape)
    tf.vertical_anchor = MSO_ANCHOR.TOP
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run(); r.text = text
    set_run(r, size=size, bold=bold, color=color, font=font)
    return shape


def add_box(slide, x, y, w, h, text, fill, line, size=14, bold=False, color=INK, align=PP_ALIGN.CENTER):
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid(); sh.fill.fore_color.rgb = fill
    sh.line.color.rgb = line; sh.line.width = Pt(1.2)
    tf = clear_frame(sh); tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]; p.alignment = align
    r = p.add_run(); r.text = text
    set_run(r, size=size, bold=bold, color=color)
    return sh


def add_arrow(slide, x1, y1, x2, y2, color=RED):
    line = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    line.line.color.rgb = color; line.line.width = Pt(1.8)
    line.line.end_arrowhead = True
    return line


def add_code_panel(slide, x, y, w, h, heading, code, accent=RED):
    panel = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    panel.fill.solid(); panel.fill.fore_color.rgb = DARK
    panel.line.color.rgb = DARK
    add_text(slide, x+0.16, y+0.10, w-0.32, 0.28, heading, size=11, bold=True, color=WHITE, font="Arial")
    tb = slide.shapes.add_textbox(Inches(x+0.16), Inches(y+0.43), Inches(w-0.30), Inches(h-0.52))
    tf = clear_frame(tb); tf.margin_left = 0; tf.margin_right = 0; tf.margin_top = 0; tf.margin_bottom = 0
    for i, line in enumerate(code.splitlines()):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(0); p.alignment = PP_ALIGN.LEFT
        r = p.add_run(); r.text = line
        set_run(r, size=9.5, color=WHITE, font="Consolas")
    return panel


def build():
    if not TEMPLATE.exists():
        raise FileNotFoundError(TEMPLATE)
    old = ROOT / "jinma-award-statement.pptx"
    if old.exists() and not BACKUP.exists():
        shutil.copy2(old, BACKUP)

    prs = Presentation(str(TEMPLATE))
    if len(prs.slides) != 8:
        raise RuntimeError(f"template must have 8 slides, got {len(prs.slides)}")

    # 01 Cover: preserve template composition; replace only the variable fields.
    group = prs.slides[0].shapes[3]
    replace_all_runs(group.shapes[1], {"XXX": "CLIProxyAPI 企业级 AI 网关建设"})
    replace_all_runs(group.shapes[2], {"XXX": "吴强辉"})

    # 02 Contents.
    s = prs.slides[1]
    set_plain(s.shapes[0], "个人概况", size=20, bold=True, color=INK)
    set_plain(s.shapes[4], "申报亮点", size=20, bold=True, color=INK)
    set_plain(s.shapes[7], "设计模式\n与代码举证", size=19, bold=True, color=INK)
    set_plain(s.shapes[10], "总结", size=20, bold=True, color=INK)

    # 03 Profile. Keep the template's arrow bullets and replace fields in place.
    s = prs.slides[2]
    replace_paragraphs(s.shapes[1], [
        "姓名：吴强辉",
        "所在部门：研发中心四部",
        "当前职级：P2｜C++开发工程师",
        "所属模块：AI 网关管理中心",
        "代码仓库：CLIProxyAPI",
        "交付状态：已完成运行验证",
    ], size=20)
    # No personal portrait was supplied; retain the template's required placeholder without fabricating an image.
    photo = s.shapes[2]
    photo.text_frame.clear()
    photo.fill.solid(); photo.fill.fore_color.rgb = WHITE
    photo.line.color.rgb = INK; photo.line.width = Pt(1)
    tf = photo.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = "个人近期照\n（正面、清晰）"
    set_run(r, size=19, color=INK)

    # 04 Highlights: tailor the template's award criteria to the project.
    s = prs.slides[3]
    box = s.shapes[1]
    tf = clear_frame(box)
    tf.margin_left = Inches(0.16); tf.margin_right = Inches(0.12); tf.margin_top = Inches(0.08); tf.margin_bottom = Inches(0.04)
    items = [
        ("算法与数据结构：", "以 usage_events 保留原始事实，Rollup 作为可重建加速层；Analytics 按 Tab 按需返回，明细采用 keyset 分页。"),
        ("架构与设计模式：", "保持 CLIProxyAPI 统一接入，Management Center 负责企业治理，Usage Service 负责采集、聚合与协调；插件通过宿主认证和路径白名单隔离。"),
        ("业务实现质量：", "企业 Key 支持导入预览、冲突处理和历史追踪；额度治理区分度量、规则、动作与恢复；上游状态不确定时保持现状。"),
        ("价值：", "将分散的 Key、额度、用量与审计能力形成统一管理闭环，支撑企业化运维和可验证交付。"),
        ("创新性：", "将 AI 用于跨文件检索、需求拆解、风险识别与测试生成，结合工程师复核形成增量交付流程。"),
        ("规范性：", "遵循现有服务边界，原始数据可重建、明细数据脱敏，测试覆盖异常、权限、并发、构建与运行场景。"),
    ]
    for label, body in items:
        add_rich_paragraph(tf, label, body, size=12.3, after=2)

    # 05 Design pattern: replace the template text area with the architecture and AI development method.
    s = prs.slides[4]
    banner = s.shapes[1]
    banner.fill.solid(); banner.fill.fore_color.rgb = PALE
    banner.line.color.rgb = RED; banner.line.width = Pt(1.2)
    set_plain(banner, "设计思想：边界清晰、能力解耦、事实可重建、扩展可控", size=19, bold=True, color=INK, align=PP_ALIGN.CENTER)
    add_box(s, 1.25, 2.20, 1.65, 0.68, "客户端 / Agent", PALE, RED, size=14, bold=True)
    add_box(s, 3.35, 2.20, 1.75, 0.68, "CLIProxyAPI\n统一接入", RED, RED, size=14, bold=True, color=WHITE)
    add_box(s, 5.60, 2.20, 2.10, 0.68, "Management Center\n企业治理", BLUE, BLUE, size=14, bold=True, color=WHITE)
    add_box(s, 5.60, 3.36, 2.10, 0.68, "Usage Service\n采集 · 聚合 · 协调", PALE, BLUE, size=13.5, bold=True)
    add_box(s, 8.35, 2.20, 2.00, 0.68, "插件 / 审计\n受控扩展", PALE, RED, size=13.5, bold=True)
    add_arrow(s, 2.90, 2.54, 3.35, 2.54, RED)
    add_arrow(s, 5.10, 2.54, 5.60, 2.54, BLUE)
    add_arrow(s, 6.65, 2.88, 6.65, 3.36, BLUE)
    add_arrow(s, 7.70, 2.54, 8.35, 2.54, RED)
    add_text(s, 1.25, 4.38, 3.45, 0.65, "AI 辅助方式\n检索调用链、归纳差异、生成候选方案", size=14, bold=True, color=INK)
    add_text(s, 4.78, 4.38, 3.45, 0.65, "工程判断\n确定边界、兼容策略、失败语义与验收标准", size=14, bold=True, color=INK)
    add_text(s, 8.30, 4.38, 2.85, 0.65, "交付原则\n增量实现、可测试、可回溯、可发布", size=14, bold=True, color=INK)

    # 06 Code evidence: pair real implementation fragments with their engineering meaning.
    s = prs.slides[5]
    panel = s.shapes[1]
    panel.fill.solid(); panel.fill.fore_color.rgb = DARK
    panel.line.color.rgb = DARK
    panel.text_frame.clear()
    add_code_panel(s, 1.30, 1.57, 5.10, 2.00, "举证 1｜企业 Key 领域模型", "interface EnterpriseKeyIdentity {\n  apiKeyHash: string;\n  userName: string;\n  departmentId?: string;\n  source: string;\n  importedAt: string;\n}", RED)
    add_code_panel(s, 6.72, 1.57, 5.10, 2.00, "举证 2｜额度规则分层", "func (c SpendLimitConfig) LimitForKey(hash string) SpendLimit {\n  if limit, ok := c.OverrideForKey(hash); ok {\n    return limit\n  }\n  return c.DefaultLimit()\n}", BLUE)
    add_text(s, 1.30, 3.95, 5.10, 0.78, "身份模型以 apiKeyHash 建立稳定关联；页面只展示非秘密元数据，原始凭证与展示数据分层管理。", size=13.2, color=INK)
    add_text(s, 6.72, 3.95, 5.10, 0.78, "Key override 优先于默认限额；费用计算、动作执行和失败请求排除分别验证，避免治理规则隐式耦合。", size=13.2, color=INK)
    add_text(s, 1.30, 5.30, 10.5, 0.28, "代码来源：Management Center 企业 Key / Usage Service spend-limit；完整实现、测试与边界说明纳入交付材料。", size=10.5, color=MUTED)

    # 07 Summary.
    s = prs.slides[6]
    summary = s.shapes[1]
    tf = clear_frame(summary)
    tf.margin_left = Inches(0.25); tf.margin_right = Inches(0.20); tf.margin_top = Inches(0.08)
    add_rich_paragraph(tf, "建设成果：", "在 CLIProxyAPI 统一接入基础上，形成覆盖企业身份、额度治理、用量分析、安全审计与工程交付的管理能力。", size=18, after=12)
    add_rich_paragraph(tf, "方法沉淀：", "以需求拆解、依赖排序、边界验证和发布门禁构建 AI 辅助开发闭环；AI 负责检索与候选生成，工程师负责取舍与验收。", size=18, after=12)
    add_rich_paragraph(tf, "复用价值：", "保留上游兼容能力，新增模块可独立演进、可重建、可验证，为企业级 AI 应用治理提供可复用基础。", size=18, after=12)

    # 08: retain company closing slide from the supplied template.
    prs.save(str(OUTPUT))
    # Keep the historical filename usable by the user, while preserving the previous 25-page deck.
    current = ROOT / "jinma-award-statement.pptx"
    shutil.copy2(OUTPUT, current)
    print(f"created {OUTPUT}")
    print(f"updated {current}")
    print(f"archived {BACKUP}")


if __name__ == "__main__":
    build()
