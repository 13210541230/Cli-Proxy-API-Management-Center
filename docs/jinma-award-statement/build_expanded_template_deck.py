from pathlib import Path
import shutil
from PIL import Image, ImageOps
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parent
TEMPLATE = Path(r"D:/Downloads/2026年金码奖申报陈述PPT模板-XX部门-姓名.pptx")
OUTPUT = ROOT / "jinma-award-statement-expanded.pptx"
CURRENT = ROOT / "jinma-award-statement.pptx"
NAMED = ROOT / "2026年金码奖申报陈述PPT-研发中心四部-吴强辉.pptx"
ARCHIVE_8 = ROOT / "jinma-award-statement-8page-template.pptx"
PHOTO = Path(r"D:/Pictures/吴强辉.jpg")
PHOTO_ASSET = ROOT / "media" / "profile-wuqh.png"

RED = RGBColor(0xED, 0x00, 0x16)
BLUE = RGBColor(0x06, 0x67, 0xB1)
INK = RGBColor(0x22, 0x22, 0x22)
MUTED = RGBColor(0x66, 0x66, 0x66)
PALE = RGBColor(0xF5, 0xF7, 0xF9)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
DARK = RGBColor(0x18, 0x23, 0x2D)
LIGHT_RED = RGBColor(0xFF, 0xF0, 0xF1)
LIGHT_BLUE = RGBColor(0xEE, 0xF6, 0xFC)


def set_run(run, size=18, bold=False, color=INK, font="Microsoft YaHei", italic=False):
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color


def clear_frame(shape):
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(0.08)
    tf.margin_right = Inches(0.08)
    tf.margin_top = Inches(0.04)
    tf.margin_bottom = Inches(0.04)
    return tf


def clear_slide(slide):
    for shape in list(slide.shapes):
        slide.shapes._spTree.remove(shape._element)


def set_plain(shape, text, size=18, bold=False, color=INK, align=None, font="Microsoft YaHei"):
    tf = clear_frame(shape)
    p = tf.paragraphs[0]
    p.alignment = align if align is not None else PP_ALIGN.LEFT
    r = p.add_run(); r.text = text
    set_run(r, size=size, bold=bold, color=color, font=font)
    return tf


def replace_all_runs(shape, replacements):
    if not hasattr(shape, "text_frame"):
        return
    for p in shape.text_frame.paragraphs:
        for r in p.runs:
            for old, new in replacements.items():
                if old in r.text:
                    r.text = r.text.replace(old, new)


def add_text(slide, x, y, w, h, text, size=16, bold=False, color=INK,
             align=PP_ALIGN.LEFT, font="Microsoft YaHei", italic=False):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = clear_frame(shape)
    tf.vertical_anchor = MSO_ANCHOR.TOP
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run(); r.text = text
    set_run(r, size=size, bold=bold, color=color, font=font, italic=italic)
    return shape


def add_page_title(slide, title, subtitle=None):
    add_text(slide, 0.33, 0.36, 12.0, 0.52, title, size=24, bold=True, color=MUTED)
    if subtitle:
        add_text(slide, 0.72, 0.94, 11.8, 0.34, subtitle, size=11.5, color=MUTED)


def add_box(slide, x, y, w, h, text, fill, line, size=14, bold=False,
            color=INK, align=PP_ALIGN.CENTER, radius=True):
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    sh = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid(); sh.fill.fore_color.rgb = fill
    sh.line.color.rgb = line; sh.line.width = Pt(1.1)
    tf = clear_frame(sh); tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]; p.alignment = align
    r = p.add_run(); r.text = text
    set_run(r, size=size, bold=bold, color=color)
    return sh


def add_arrow(slide, x1, y1, x2, y2, color=RED, width=1.8):
    line = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    line.line.color.rgb = color; line.line.width = Pt(width); line.line.end_arrowhead = True
    return line


def add_rich_paragraph(tf, label, body, size=15, color=INK, after=4):
    # Reuse the initial paragraph after clear_frame; do not leave an empty
    # leading paragraph that creates unexplained vertical whitespace.
    if len(tf.paragraphs) == 1 and not tf.paragraphs[0].text:
        p = tf.paragraphs[0]
    else:
        p = tf.add_paragraph()
    p.alignment = PP_ALIGN.LEFT
    p.space_after = Pt(after)
    p.line_spacing = 1.08
    r1 = p.add_run(); r1.text = label; set_run(r1, size=size, bold=True, color=color)
    r2 = p.add_run(); r2.text = body; set_run(r2, size=size, color=color)
    return p


def add_labeled_bullets(slide, x, y, w, h, items, size=18, after=9):
    """Template-like profile list with a stable label/value hierarchy."""
    sh = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = clear_frame(sh)
    tf.margin_left = Inches(0.02); tf.margin_right = Inches(0.02)
    tf.margin_top = Inches(0.02); tf.margin_bottom = Inches(0.02)
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.level = 0; p.space_after = Pt(after); p.line_spacing = 1.05
        if "：" in item:
            label, value = item.split("：", 1)
        else:
            label, value = item, ""
        bullet = p.add_run(); bullet.text = "• "; set_run(bullet, size=size, bold=True, color=RED)
        r1 = p.add_run(); r1.text = label + ("：" if value else ""); set_run(r1, size=size, bold=True, color=INK)
        if value:
            r2 = p.add_run(); r2.text = value; set_run(r2, size=size, color=INK)
    return sh


def add_card(slide, x, y, w, h, title, body, accent=RED, fill=PALE,
             title_size=14, body_size=12.2):
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid(); sh.fill.fore_color.rgb = fill
    sh.line.color.rgb = accent; sh.line.width = Pt(1.0)
    tf = clear_frame(sh)
    tf.margin_left = Inches(0.16); tf.margin_right = Inches(0.14)
    tf.margin_top = Inches(0.12); tf.margin_bottom = Inches(0.08)
    p = tf.paragraphs[0]; p.space_after = Pt(6)
    r = p.add_run(); r.text = title; set_run(r, size=title_size, bold=True, color=accent)
    p2 = tf.add_paragraph(); p2.space_after = Pt(0)
    r2 = p2.add_run(); r2.text = body; set_run(r2, size=body_size, color=INK)
    return sh


def add_bullet_list(slide, x, y, w, h, items, size=14, color=INK, accent=RED):
    sh = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = clear_frame(sh); tf.margin_left = Inches(0.02); tf.margin_right = Inches(0.02)
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.level = 0; p.space_after = Pt(8)
        p.text = "• " + item
        for r in p.runs: set_run(r, size=size, color=color)
    return sh


def add_code_panel(slide, x, y, w, h, heading, code, accent=RED, code_size=9.0):
    panel = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    panel.fill.solid(); panel.fill.fore_color.rgb = DARK; panel.line.color.rgb = DARK
    add_text(slide, x+0.16, y+0.10, w-0.3, 0.28, heading, size=11, bold=True, color=WHITE, font="Arial")
    tb = slide.shapes.add_textbox(Inches(x+0.16), Inches(y+0.43), Inches(w-0.30), Inches(h-0.52))
    tf = clear_frame(tb); tf.margin_left = 0; tf.margin_right = 0; tf.margin_top = 0; tf.margin_bottom = 0
    for i, line in enumerate(code.splitlines()):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph(); p.space_after = Pt(0)
        r = p.add_run(); r.text = line; set_run(r, size=code_size, color=WHITE, font="Consolas")
    return panel


def add_picture(slide, path, x, y, w, h, line=BLUE):
    pic = slide.shapes.add_picture(str(path), Inches(x), Inches(y), width=Inches(w), height=Inches(h))
    frame = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    frame.fill.background(); frame.line.color.rgb = line; frame.line.width = Pt(1.0)
    return pic


def add_flow(slide, labels, x=0.8, y=2.05, box_w=2.2, box_h=0.72, gap=0.42):
    for i, label in enumerate(labels):
        bx = x + i * (box_w + gap)
        fill, line, color = (RED, RED, WHITE) if i == 1 else (PALE, BLUE, INK)
        add_box(slide, bx, y, box_w, box_h, label, fill, line, size=13, bold=True, color=color)
        if i < len(labels)-1:
            add_arrow(slide, bx+box_w, y+box_h/2, bx+box_w+gap, y+box_h/2, BLUE)


def new_inner_slide(prs, title=None, subtitle=None):
    # Titles are written by each page builder so existing template slides and
    # appended slides share exactly the same rendering path.
    # Layout 2 carries the template's title-bar red square.  The former
    # section-title layout omitted it, so appended pages from page 8 onward
    # looked different from the template-based opening pages.
    slide = prs.slides.add_slide(prs.slide_layouts[2])
    clear_slide(slide)
    return slide


def remove_slide(prs, index):
    sld_id = prs.slides._sldIdLst[index]
    prs.slides._sldIdLst.remove(sld_id)


def prepare_profile_photo():
    if not PHOTO.exists():
        raise FileNotFoundError(PHOTO)
    PHOTO_ASSET.parent.mkdir(parents=True, exist_ok=True)
    target = (770, 930)
    image = Image.open(PHOTO).convert("RGB")
    image = ImageOps.fit(image, target, method=Image.Resampling.LANCZOS, centering=(0.5, 0.42))
    # Keep the photo rectangular so the image edge stays exactly inside the
    # square-corner frame on the profile page.
    image.save(PHOTO_ASSET)
    return PHOTO_ASSET


def populate_profile(slide):
    clear_slide(slide)
    add_page_title(slide, "自我介绍")
    rows = [
        ("姓名", "吴强辉"),
        ("所在部门", "研发中心四部"),
        ("当前职级", "P2｜C++开发工程师"),
        ("所属模块", "CLIProxyAPI"),
        ("代码仓库", "https://github.com/13210541230/CLIProxyAPI"),
        ("交付状态", "已完成运行验证"),
    ]
    start_y = 1.48
    for i, (label, value) in enumerate(rows):
        y = start_y + i * 0.68
        # The template uses a small black arrow marker, a bold field label,
        # and a regular-weight value on one shared baseline.
        marker = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(1.15), Inches(y + 0.10), Inches(0.18), Inches(0.16))
        marker.fill.solid(); marker.fill.fore_color.rgb = INK
        marker.line.color.rgb = INK
        add_text(slide, 1.48, y, 1.42, 0.40, label + "：", size=18, bold=True, color=INK)
        value_size = 11.5 if label == "代码仓库" else 18
        add_text(slide, 2.86, y, 4.35, 0.40, value, size=value_size, color=INK)

    photo_asset = prepare_profile_photo()
    photo_x, photo_y, photo_w, photo_h = 7.65, 1.50, 3.85, 4.65
    slide.shapes.add_picture(str(photo_asset), Inches(photo_x), Inches(photo_y), width=Inches(photo_w), height=Inches(photo_h))
    frame = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(photo_x), Inches(photo_y), Inches(photo_w), Inches(photo_h))
    frame.fill.background()
    frame.line.color.rgb = INK; frame.line.width = Pt(1.0)


def populate_highlights(slide):
    clear_slide(slide)
    add_page_title(slide, "申报亮点", "围绕技术先进性、工程质量与业务价值，形成可验证的企业级 AI 网关建设成果。")
    items = [
        ("01", "企业级治理能力建设", "在 CLIProxyAPI 统一接入基础上，构建企业身份、额度治理、用量分析与安全审计能力，形成覆盖配置、运行、留痕与交付的管理闭环。", RED, LIGHT_RED),
        ("02", "AI 辅助研发与工程化交付", "以 AI 辅助跨文件检索、需求拆解、风险识别与测试生成，结合事实可重建、失败可解释、扩展受控与发布门禁，提升研发交付的可复核性与复用价值。", BLUE, LIGHT_BLUE),
    ]
    for num, title, body, accent, fill, in items:
        y = 1.70 if num == "01" else 3.95
        badge = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.92), Inches(y), Inches(0.82), Inches(0.82))
        badge.fill.solid(); badge.fill.fore_color.rgb = accent
        badge.line.color.rgb = accent
        tf_badge = clear_frame(badge); tf_badge.vertical_anchor = MSO_ANCHOR.MIDDLE
        p_badge = tf_badge.paragraphs[0]; p_badge.alignment = PP_ALIGN.CENTER
        r_badge = p_badge.add_run(); r_badge.text = num; set_run(r_badge, size=20, bold=True, color=WHITE)
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.98), Inches(y), Inches(10.35), Inches(1.55))
        card.fill.solid(); card.fill.fore_color.rgb = fill
        card.line.color.rgb = accent; card.line.width = Pt(1.0)
        tf = clear_frame(card)
        tf.margin_left = Inches(0.28); tf.margin_right = Inches(0.24)
        tf.margin_top = Inches(0.14); tf.margin_bottom = Inches(0.10)
        p1 = tf.paragraphs[0]; p1.alignment = PP_ALIGN.LEFT; p1.space_after = Pt(6)
        r1 = p1.add_run(); r1.text = title; set_run(r1, size=17, bold=True, color=accent)
        p2 = tf.add_paragraph(); p2.alignment = PP_ALIGN.LEFT; p2.line_spacing = 1.08
        r2 = p2.add_run(); r2.text = body; set_run(r2, size=13.2, color=INK)


def populate_needs(slide):
    clear_slide(slide)
    add_page_title(slide, "建设背景｜企业需求驱动能力扩展", "从统一接入出发，面向企业运营补齐身份、治理、分析、安全与交付能力。")
    card_xs = [0.82, 3.79, 6.76, 9.73]
    card_w = 2.75
    add_card(slide, card_xs[0], 1.55, card_w, 2.08, "统一身份", "将 Key 与使用人、部门、来源和历史关系关联，形成可审计的企业身份模型。", RED, LIGHT_RED)
    add_card(slide, card_xs[1], 1.55, card_w, 2.08, "额度治理", "统一 cost / tokens 度量，支持默认规则、Key 覆盖、暂停、降级、通知与恢复。", BLUE, LIGHT_BLUE)
    add_card(slide, card_xs[2], 1.55, card_w, 2.08, "用量分析", "在高并发访问下保留原始事实，构建可重建聚合和可解释的明细查询链路。", RED, LIGHT_RED)
    add_card(slide, card_xs[3], 1.55, card_w, 2.08, "安全交付", "插件扩展具备来源、方法、路径和凭证边界，核心管理能力保持稳定可用。", BLUE, LIGHT_BLUE)
    add_flow(slide, ["统一接入", "企业治理", "留痕与分析", "安全交付"], x=1.25, y=4.45, box_w=2.32, gap=0.54)
    add_text(slide, 1.05, 5.55, 11.1, 0.4, "建设目标：将分散的接入、身份、额度、用量与审计能力组织为可持续演进的企业治理底座。", size=15, bold=True, color=INK, align=PP_ALIGN.CENTER)


def populate_architecture(slide):
    clear_slide(slide)
    add_page_title(slide, "总体架构与设计模式｜边界清晰、能力解耦、事实可重建")
    add_box(slide, 0.82, 1.65, 1.85, 0.72, "客户端 / Agent", PALE, RED, size=14, bold=True)
    add_box(slide, 3.05, 1.65, 2.05, 0.72, "CLIProxyAPI\n统一接入", RED, RED, size=14, bold=True, color=WHITE)
    add_box(slide, 5.55, 1.65, 2.25, 0.72, "Management Center\n企业治理", BLUE, BLUE, size=14, bold=True, color=WHITE)
    add_box(slide, 5.55, 3.00, 2.25, 0.72, "Usage Service\n采集 · 聚合 · 协调", PALE, BLUE, size=13.5, bold=True)
    add_box(slide, 8.35, 1.65, 2.20, 0.72, "插件 / 审计\n受控扩展", PALE, RED, size=13.5, bold=True)
    add_arrow(slide, 2.67, 2.01, 3.05, 2.01, RED)
    add_arrow(slide, 5.10, 2.01, 5.55, 2.01, BLUE)
    add_arrow(slide, 6.68, 2.37, 6.68, 3.00, BLUE)
    add_arrow(slide, 7.80, 2.01, 8.35, 2.01, RED)
    add_card(slide, 0.90, 4.35, 3.55, 1.25, "分层职责", "接入层协议兼容；管理层承载企业规则；Usage Service 负责事实采集与聚合。", RED, PALE, body_size=12)
    add_card(slide, 4.85, 4.35, 3.55, 1.25, "数据策略", "usage_events 保存原始事实；Rollup 可删除、可重建；覆盖不足时回退 raw 并说明完整性。", BLUE, PALE, body_size=12)
    add_card(slide, 8.80, 4.35, 3.55, 1.25, "扩展策略", "插件通过宿主认证、来源校验、方法限制和路径白名单进入受控边界。", RED, PALE, body_size=12)


def populate_ai_method(slide):
    clear_slide(slide)
    add_page_title(slide, "AI辅助开发方法｜将业务目标转化为可执行工作包", "AI 负责信息组织与候选生成，工程师负责边界判断、取舍和最终验收。")
    add_flow(slide, ["需求原文", "AI检索与拆解", "任务与依赖", "实现与验收"], x=0.65, y=1.60, box_w=2.55, gap=0.45)
    add_card(slide, 0.82, 3.25, 3.65, 1.80, "信息组织", "跨文件检索接口、数据表、调用链与现有约束；归纳候选影响面，避免只改表面页面。", RED, LIGHT_RED)
    add_card(slide, 4.84, 3.25, 3.65, 1.80, "工程判断", "明确兼容边界、失败语义、数据脱敏、状态转换和取舍理由，将风险转化为验收条款。", BLUE, LIGHT_BLUE)
    add_card(slide, 8.86, 3.25, 3.65, 1.80, "交付物", "输出任务清单、依赖关系、测试命令、构建结果、运行验证和可回溯的变更说明。", RED, LIGHT_RED)
    add_text(slide, 0.95, 5.65, 11.3, 0.42, "适用范围：企业 Key、额度治理、Usage Analytics、插件审计及其跨服务协同。", size=13, color=MUTED)


def populate_plan(slide):
    clear_slide(slide)
    add_page_title(slide, "执行计划与验收矩阵｜按依赖推进、按证据收口", "每项建设经历领域建模、接口实现、边界验证和运行交付四个阶段。")
    steps = [
        ("01 · 领域建模", "明确身份、额度、事实源、状态和安全边界。", "数据模型与需求矩阵"),
        ("02 · 接口实现", "按现有服务边界拆分 API、存储、页面和迁移任务。", "可运行增量实现"),
        ("03 · 边界验证", "覆盖失败、重试、恢复、权限、性能和脱敏场景。", "测试矩阵与结果"),
        ("04 · 运行交付", "完成构建、启动、内嵌资源、冒烟和视觉一致性验证。", "发布包与复核记录"),
    ]
    xs = [0.78, 3.85, 6.92, 9.99]
    for i, (head, body, out) in enumerate(steps):
        if i < 3: add_arrow(slide, xs[i]+2.25, 2.20, xs[i+1], 2.20, RED)
        add_box(slide, xs[i], 1.85, 2.25, 0.70, head, RED, RED, size=13, bold=True, color=WHITE)
        add_text(slide, xs[i], 2.85, 2.45, 0.95, body, size=12, color=MUTED)
        add_text(slide, xs[i], 3.78, 2.45, 0.48, "产物：" + out, size=11.5, color=INK)
    add_text(slide, 0.82, 5.00, 2.65, 0.3, "ACCEPTANCE MATRIX", size=13, bold=True, color=BLUE, font="Arial")
    add_text(slide, 0.82, 5.45, 11.5, 0.45, "功能正确性｜数据完整性｜安全边界｜并发性能｜构建完整性｜运行可用性｜视觉一致性", size=14, color=INK)


def add_module_chips(slide, modules, y=5.42):
    """Show the module scope as a deliberate template-style footer row."""
    if isinstance(modules, str):
        labels = [item.strip() for item in modules.split("、") if item.strip()]
    else:
        labels = list(modules)
    add_text(slide, 0.82, y + 0.03, 1.02, 0.28, "相关模块", size=10.5, bold=True, color=BLUE)
    if not labels:
        return
    x = 1.92
    gap = 0.12
    available = 10.55
    width = (available - gap * (len(labels) - 1)) / len(labels)
    for label in labels:
        add_box(slide, x, y, width, 0.38, label, PALE, BLUE, size=9.5, bold=False, color=INK, radius=True)
        x += width + gap


def requirement_page(slide, number, title, subtitle, cards, flow_text, modules):
    clear_slide(slide)
    add_page_title(slide, title, subtitle)
    for i, (head, body, accent) in enumerate(cards):
        x = 0.82 + i * 4.02
        add_card(slide, x, 1.55, 3.55, 2.18, head, body, accent, LIGHT_RED if accent == RED else LIGHT_BLUE, title_size=14, body_size=12.2)
    add_text(slide, 1.0, 4.35, 11.3, 0.5, flow_text, size=15, bold=True, color=INK, align=PP_ALIGN.CENTER)
    add_module_chips(slide, modules, y=5.42)


def populate_key_solution(slide):
    add_page_title(slide, "企业 Key｜实现路径", "将批量导入从一次性写入转化为可预览、可修正、可追踪的治理流程。")
    add_flow(slide, ["上游 Key", "预览与规范化", "身份绑定", "历史追踪"], x=0.80, y=1.52, box_w=2.35, gap=0.42)
    add_card(slide, 0.85, 3.10, 3.35, 1.65, "输入处理", "解析 CSV，保留原始 Key 与非秘密元数据，输出可确认的导入预览。", RED, LIGHT_RED)
    add_card(slide, 4.45, 3.10, 3.35, 1.65, "身份映射", "以 apiKeyHash 建立稳定关联，绑定使用人、部门、来源和导入时间，不改变原有接入方式。", BLUE, LIGHT_BLUE)
    add_card(slide, 8.05, 3.10, 3.35, 1.65, "结果可追溯", "成功、警告和错误可区分；冲突和部门缺失可修正；导入结果可复核。", RED, LIGHT_RED)
    add_text(slide, 0.90, 5.35, 11.3, 0.42, "实现结果：企业 Key 从“凭证集合”升级为可治理、可检索、可审计的身份对象。", size=14, bold=True, color=INK, align=PP_ALIGN.CENTER)


def populate_key_code(slide):
    add_page_title(slide, "企业 Key｜实现结构与关键逻辑", "从目录边界、领域模型到展示隔离，说明企业身份能力如何作为旁路治理增量接入。")
    code = "type EnterpriseKeyIdentity struct {\n  APIKeyHash   string\n  UserName     string\n  DepartmentID *string\n  Source       string\n  ImportedAt   time.Time\n}\n\nfunc bindIdentity(row ImportRow) Identity {\n  return Identity{APIKeyHash: hash(row.APIKey),\n    UserName: row.UserName, DepartmentID: row.DepartmentID,\n    Source: row.Source, ImportedAt: time.Now()}\n}"
    tree = "enterprise-key/\n├─ model/identity.go\n├─ import/preview.go\n├─ binding/service.go\n├─ history/repository.go\n└─ api/enterprise_keys.go"
    add_code_panel(slide, 0.82, 1.48, 6.15, 3.62, "领域模型｜身份关联与写入边界", code, RED, 8.8)
    add_code_panel(slide, 7.22, 1.48, 5.05, 2.28, "代码结构｜模块职责", tree, BLUE, 9.8)
    add_card(slide, 7.22, 4.02, 5.05, 1.18, "安全展示", "原始凭证只参与内部关联；页面只返回掩码、Hash 和非秘密治理元数据。", BLUE, LIGHT_BLUE, body_size=11.8)
    add_card(slide, 0.82, 5.35, 2.92, 1.08, "稳定关联", "apiKeyHash 支撑跨导入、跨页面和历史查询。", RED, LIGHT_RED, body_size=11.5)
    add_card(slide, 3.98, 5.35, 2.99, 1.08, "兼容策略", "保留 CPA Key 结构，治理字段旁路扩展。", RED, LIGHT_RED, body_size=11.5)


def populate_key_accept(slide):
    add_page_title(slide, "企业 Key｜验收结果")
    add_card(slide, 0.82, 1.55, 3.55, 1.82, "成功结果", "批量导入完成后，Key、使用人、部门和来源关系可查询；导入结果可复核。", RED, LIGHT_RED, body_size=11.7)
    add_card(slide, 4.84, 1.55, 3.55, 1.82, "警告与错误", "重复 Key、部门缺失和格式错误分别反馈；可修正项不阻断其余有效记录。", BLUE, LIGHT_BLUE, body_size=11.7)
    add_card(slide, 8.86, 1.55, 3.55, 1.82, "安全展示", "页面只展示安全版本的凭证信息，原始 Key 不进入展示面；历史记录保留来源和时间。", RED, LIGHT_RED, body_size=11.7)
    add_text(slide, 0.90, 3.78, 11.2, 0.42, "验收结论：导入流程具备可预览、可修正、可追踪和可审计能力。", size=14.5, bold=True, color=INK, align=PP_ALIGN.CENTER)
    add_picture(slide, ROOT / "media" / "account-overview.png", 4.20, 4.48, 4.95, 2.60, line=BLUE)


def populate_quota_solution(slide):
    add_page_title(slide, "额度治理｜实现路径", "将额度规则、Key 覆盖、超限动作和上游状态纳入同一治理链路。")
    add_flow(slide, ["用量事实", "成本 / Token 计算", "规则匹配", "治理动作"], x=0.82, y=1.55, box_w=2.35, gap=0.42)
    add_card(slide, 0.82, 3.12, 3.55, 1.55, "统一度量", "支持 cost / tokens 两种模式，统一日、周窗口与失败请求处理口径。", RED, LIGHT_RED)
    add_card(slide, 4.84, 3.12, 3.55, 1.55, "分层策略", "默认限额优先，Key override 覆盖；超限动作包含 pause、downgrade 与 notify。", BLUE, LIGHT_BLUE)
    add_card(slide, 8.86, 3.12, 3.55, 1.55, "状态协调", "暂停、降级、通知和恢复保持 reason、expires_at 等可解释状态。", RED, LIGHT_RED)
    add_text(slide, 0.88, 5.32, 11.2, 0.42, "实现结果：额度治理从静态配置升级为可度量、可执行、可恢复的运行状态协调。", size=14, bold=True, color=INK, align=PP_ALIGN.CENTER)


def populate_quota_code(slide):
    add_page_title(slide, "额度治理｜实现结构与关键逻辑", "规则、动作与上游状态分别建模，确保额度治理可计算、可执行、可恢复。")
    code1 = "type SpendLimitConfig struct {\n  Default   SpendLimit\n  Overrides map[string]SpendLimit\n}\n\nfunc (c SpendLimitConfig) LimitForKey(hash string) SpendLimit {\n  if limit, ok := c.Overrides[hash]; ok {\n    return limit\n  }\n  return c.Default\n}"
    code2 = "func reconcile(current State, upstream *Status) State {\n  if upstream == nil {\n    return current.WithReason(\"upstream_unavailable\")\n  }\n  if upstream.Paused {\n    return current.Pause(upstream.ExpiresAt)\n  }\n  return current.Recover()\n}"
    add_code_panel(slide, 0.82, 1.48, 5.75, 3.16, "规则层｜Key 覆盖优先于默认限额", code1, RED, 8.8)
    add_code_panel(slide, 6.78, 1.48, 5.75, 3.16, "状态层｜不可确认时保持现状", code2, BLUE, 8.8)
    add_card(slide, 0.82, 5.02, 5.75, 1.12, "分层价值", "费用计算、规则匹配与治理动作解耦，便于分别测试、回溯和替换。", RED, LIGHT_RED, body_size=11.8)
    add_card(slide, 6.78, 5.02, 5.75, 1.12, "异常边界", "502/503/504/429 或网络不可确认时，不执行猜测性暂停或恢复。", BLUE, LIGHT_BLUE, body_size=11.8)


def populate_quota_accept(slide):
    add_page_title(slide, "额度治理｜验收结果")
    add_card(slide, 0.82, 1.55, 3.55, 1.82, "计算验收", "分别验证 cost / tokens、日 / 周窗口、缓存 Token 和失败请求排除，结果与运营口径一致。", RED, LIGHT_RED, body_size=11.7)
    add_card(slide, 4.84, 1.55, 3.55, 1.82, "动作验收", "默认规则与 Key 覆盖、pause、降级、通知和恢复窗口分别验证，动作可追踪。", BLUE, LIGHT_BLUE, body_size=11.7)
    add_card(slide, 8.86, 1.55, 3.55, 1.82, "异常验收", "上游状态不可确认时保留现状；重试次数和退避策略可验证，避免误治理。", RED, LIGHT_RED, body_size=11.7)
    add_text(slide, 0.90, 3.78, 11.2, 0.42, "验收结论：额度治理具备统一口径、明确动作、异常安全和状态恢复能力。", size=14.5, bold=True, color=INK, align=PP_ALIGN.CENTER)
    add_picture(slide, ROOT / "media" / "quota-management.png", 0.82, 4.48, 5.75, 2.60, line=BLUE)
    add_picture(slide, ROOT / "media" / "credential-quota.png", 6.78, 4.48, 5.75, 2.60, line=RED)


def populate_analytics_solution(slide):
    add_page_title(slide, "用量分析｜实现路径", "以可重建数据链路支撑高并发查询，同时保持来源与完整性可解释。")
    add_flow(slide, ["usage_events", "Hourly / Daily Rollup", "Tab 按需返回", "Keyset 明细"], x=0.56, y=1.55, box_w=2.62, gap=0.26)
    add_card(slide, 0.82, 3.12, 3.55, 1.55, "原始事实", "usage_events 保留原始事件，作为重建聚合结果和核对查询的事实源。", RED, LIGHT_RED)
    add_card(slide, 4.84, 3.12, 3.55, 1.55, "聚合加速", "按小时、按日和维度生成 Rollup；进度与派生写入保持原子一致。", BLUE, LIGHT_BLUE)
    add_card(slide, 8.86, 3.12, 3.55, 1.55, "解释性查询", "筛选或部分日期未被覆盖时回退 raw，并返回 source / complete 元数据。", RED, LIGHT_RED)
    add_text(slide, 0.88, 5.32, 11.2, 0.42, "实现结果：查询性能、数据完整性与结果可解释性同时纳入设计目标。", size=14, bold=True, color=INK, align=PP_ALIGN.CENTER)


def populate_analytics_code(slide):
    add_page_title(slide, "用量分析｜实现结构与关键逻辑", "以原始事实、可重建聚合和稳定分页契约支撑高并发查询。")
    code = "type AnalyticsQuery struct {\n  Include []string\n  Cursor  *Cursor\n}\n\nfunc advanceCheckpoint(tx *Tx, next int64) error {\n  if err := tx.AggregateAndCommit(); err != nil {\n    return err\n  }\n  return tx.AdvanceCheckpoint(next)\n}\n\n// detail: ORDER BY timestamp_ms, id\n// cursor: (timestamp_ms, id)"
    tree = "usage-service/\n├─ collector/pending_items.go\n├─ rollup/hourly_daily.go\n├─ analytics/query.go\n├─ analytics/detail_keyset.go\n└─ api/monitoring.go"
    add_code_panel(slide, 0.82, 1.48, 6.35, 3.46, "数据契约｜聚合水位与明细分页", code, RED, 8.7)
    add_code_panel(slide, 7.52, 1.48, 4.75, 2.30, "代码结构｜事实到查询", tree, BLUE, 9.2)
    add_card(slide, 7.52, 4.06, 4.75, 1.08, "覆盖边界", "Rollup 未覆盖或筛选不支持时回退 raw，并返回 source / complete。", RED, LIGHT_RED, body_size=11.5)
    add_card(slide, 7.52, 5.30, 4.75, 1.08, "明细契约", "(timestamp_ms, id) keyset 分页，返回真实 total_count，避免重复与漏读。", BLUE, LIGHT_BLUE, body_size=11.5)


def populate_analytics_accept(slide):
    add_page_title(slide, "用量分析｜验收结果")
    add_card(slide, 0.82, 1.55, 3.55, 1.82, "性能验收", "按 Tab 只返回当前页面所需聚合，避免每次请求整批拉取并在内存汇总。", RED, LIGHT_RED, body_size=11.7)
    add_card(slide, 4.84, 1.55, 3.55, 1.82, "完整性验收", "覆盖水位、派生写入和 checkpoint 同步验证；失败时不推进边界。", BLUE, LIGHT_BLUE, body_size=11.7)
    add_card(slide, 8.86, 1.55, 3.55, 1.82, "解释验收", "source、complete、total_count 和筛选语义可核对，避免静默少统计。", RED, LIGHT_RED, body_size=11.7)
    add_text(slide, 0.90, 3.78, 11.2, 0.42, "验收结论：高并发读取与数据准确、来源透明、结果可重建同时满足。", size=14.5, bold=True, color=INK, align=PP_ALIGN.CENTER)
    add_picture(slide, ROOT / "media" / "request-monitoring.png", 0.82, 4.48, 5.75, 2.60, line=BLUE)
    add_picture(slide, ROOT / "media" / "request-monitoring-detail.png", 6.78, 4.48, 5.75, 2.60, line=RED)


def populate_security_solution(slide):
    add_page_title(slide, "插件审计｜实现路径", "在扩展审计能力的同时，保持核心管理面板的稳定性与独立可用。")
    add_flow(slide, ["审计需求", "安全契约", "来源与路径校验", "独立交付"], x=0.70, y=1.55, box_w=2.50, gap=0.30)
    add_card(slide, 0.82, 3.12, 3.55, 1.55, "正向能力", "企业审计插件可读取授权范围内的非秘密 metadata、auth-files 投影和模型目录。", RED, LIGHT_RED)
    add_card(slide, 4.84, 3.12, 3.55, 1.55, "拒绝路径", "校验 iframe origin、HTTP method、插件路径和资源 allowlist，拒绝越权访问。", BLUE, LIGHT_BLUE)
    add_card(slide, 8.86, 3.12, 3.55, 1.55, "兼容结果", "旧 CPA 不具备插件 API 时，核心管理面板保持可用，插件能力不阻断主流程。", RED, LIGHT_RED)
    add_text(slide, 0.88, 5.32, 11.2, 0.42, "实现结果：扩展能力与核心治理能力解耦，安全边界和兼容降级同时可验证。", size=14, bold=True, color=INK, align=PP_ALIGN.CENTER)


def populate_security_code(slide):
    add_page_title(slide, "插件审计｜实现结构与关键逻辑", "通过来源、方法、路径和凭证多重校验，将扩展能力限制在可验证范围内。")
    code = "func handleBridge(req Request) error {\n  if !allowedOrigin(req.Origin) {\n    return ErrForbidden\n  }\n  if !allowedMethod(req.Method) {\n    return ErrMethodNotAllowed\n  }\n  if !allowedResource(req.Path, allowlist) {\n    return ErrResourceDenied\n  }\n  return dispatch(req)\n}\n\n// management key never enters plugin page"
    tree = "plugin-host/\n├─ bridge/origin.go\n├─ bridge/methods.go\n├─ resources/allowlist.go\n├─ models/readonly_projection.go\n└─ audit/enterprise_access.go"
    add_code_panel(slide, 0.82, 1.48, 6.45, 3.52, "安全边界｜请求校验与凭证隔离", code, RED, 8.6)
    add_code_panel(slide, 7.62, 1.48, 4.65, 2.28, "代码结构｜插件宿主边界", tree, BLUE, 8.9)
    add_card(slide, 7.62, 4.02, 4.65, 1.08, "资源校验", "规范化路径后匹配 allowlist，拒绝路径穿越、未知 pluginId 和非法资源。", BLUE, LIGHT_BLUE, body_size=11.5)
    add_card(slide, 7.62, 5.28, 4.65, 1.08, "数据隔离", "Management Key 不进入插件页面；明细使用掩码、Hash 和非秘密 metadata。", RED, LIGHT_RED, body_size=11.5)


def populate_security_accept(slide):
    add_page_title(slide, "插件审计｜验收结果")
    add_card(slide, 0.82, 1.55, 3.55, 1.82, "正向能力", "授权范围内的审计读取、页面通信和资源加载可正常完成，插件可独立交付。", RED, LIGHT_RED, body_size=11.7)
    add_card(slide, 4.84, 1.55, 3.55, 1.82, "拒绝能力", "foreign origin、越权 method、非法路径、未知 pluginId 和缺少凭证均被拒绝。", BLUE, LIGHT_BLUE, body_size=11.7)
    add_card(slide, 8.86, 1.55, 3.55, 1.82, "兼容结果", "插件 API 不可用时核心管理能力保持可用；测试结果纳入发布验收。", RED, LIGHT_RED, body_size=11.7)
    add_text(slide, 0.90, 3.78, 11.2, 0.42, "验收结论：插件扩展与核心管理能力隔离，扩展可用性、安全边界和数据脱敏均得到验证。", size=14.0, bold=True, color=INK, align=PP_ALIGN.CENTER)
    add_picture(slide, ROOT / "media" / "inspection.png", 0.82, 4.48, 5.75, 2.60, line=BLUE)
    add_picture(slide, ROOT / "media" / "inspection-policy.png", 6.78, 4.48, 5.75, 2.60, line=RED)


def populate_summary(slide):
    clear_slide(slide)
    add_page_title(slide, "总结")
    add_card(slide, 0.82, 1.55, 3.55, 2.15, "建设成果", "基于 CLIProxyAPI 统一接入能力，形成企业身份、额度治理、用量分析、安全审计与工程交付的完整能力链。", RED, LIGHT_RED, title_size=15, body_size=12.2)
    add_card(slide, 4.84, 1.55, 3.55, 2.15, "方法沉淀", "以需求拆解、依赖排序、边界验证和发布门禁构建 AI 辅助开发闭环；AI 负责检索与候选生成，工程师负责取舍与验收。", BLUE, LIGHT_BLUE, title_size=15, body_size=12.2)
    add_card(slide, 8.86, 1.55, 3.55, 2.15, "复用价值", "保留上游兼容能力，新增模块可独立演进、可重建、可验证，为企业级 AI 应用治理提供可复用基础。", RED, LIGHT_RED, title_size=15, body_size=12.2)
    add_flow(slide, ["统一接入", "企业治理", "可运营", "可交付"], x=1.18, y=4.55, box_w=2.35, gap=0.48)
    add_text(slide, 0.95, 5.72, 11.3, 0.42, "从开源代理能力延伸至企业治理底座，完成从需求、设计、实现到验收的闭环建设。", size=14.5, bold=True, color=INK, align=PP_ALIGN.CENTER)


def build():
    if not TEMPLATE.exists():
        raise FileNotFoundError(TEMPLATE)
    if CURRENT.exists() and not ARCHIVE_8.exists():
        shutil.copy2(CURRENT, ARCHIVE_8)
    prs = Presentation(str(TEMPLATE))
    if len(prs.slides) != 8:
        raise RuntimeError(f"template must have 8 slides, got {len(prs.slides)}")

    # Keep the template cover, contents, profile, highlights, and closing page.
    cover = prs.slides[0]
    group = cover.shapes[3]
    replace_all_runs(group.shapes[1], {"XXX": "CLIProxyAPI 企业级 AI 网关建设"})
    replace_all_runs(group.shapes[2], {"XXX": "吴强辉"})

    contents = prs.slides[1]
    set_plain(contents.shapes[0], "个人概况", size=20, bold=True, color=INK)
    set_plain(contents.shapes[4], "申报亮点", size=20, bold=True, color=INK)
    set_plain(contents.shapes[7], "设计模式\n与实现结构", size=19, bold=True, color=INK)
    set_plain(contents.shapes[10], "总结", size=20, bold=True, color=INK)

    populate_profile(prs.slides[2])
    populate_highlights(prs.slides[3])
    populate_needs(prs.slides[4])

    # Reuse the template's content slides for the first section pages, then
    # append the detailed pages. The original closing slide is moved to the
    # end after all content is added, preserving the template composition.
    populate_architecture(prs.slides[5])
    populate_ai_method(prs.slides[6])
    thanks_sld_id = prs.slides._sldIdLst[7]

    # Section 03: design, method, execution and paired case evidence.
    populate_plan(new_inner_slide(prs))

    # Requirement -> solution -> code -> acceptance, repeated for four real development cases.
    requirement_page(new_inner_slide(prs, "企业 Key｜建设要求", "面向 Key 批量纳管、人员部门关联和历史审计，建立稳定的企业身份模型。"), 1, "企业 Key｜建设要求", "上游 Key 保持协议兼容，旁路建立人员、部门、来源与历史关系。", [
        ("治理目标", "Key 与使用人、部门、邮箱及来源建立关系；保持原有接入方式。", RED),
        ("需求拆解", "批量导入、CSV 预览、部门解析、冲突处理、人工修正与历史查询。", BLUE),
        ("验收标准", "成功、警告、错误可区分；结果可追溯；展示面不暴露明文凭证。", RED),
    ], "需求 → 预览与确认 → 写入与绑定 → 历史追踪", "企业 Key 管理、导入预览、绑定元数据、部门筛选与状态展示")
    populate_key_solution(new_inner_slide(prs, "企业 Key｜实现路径", "将批量导入从一次性写入转化为可预览、可修正、可追踪的治理流程。"))
    populate_key_code(new_inner_slide(prs, "企业 Key｜实现结构与关键逻辑", "身份模型承载治理关系，展示层只提供非秘密元数据。"))
    populate_key_accept(new_inner_slide(prs, "企业 Key｜验收结果"))

    requirement_page(new_inner_slide(prs, "额度治理｜建设要求", "以统一度量支撑默认规则、Key 覆盖、超限动作和上游状态安全协调。"), 2, "额度治理｜建设要求", "以 cost 或 Token 作为统一度量，将规则、动作与恢复纳入同一链路。", [
        ("度量要求", "支持 cost / tokens 两种模式，统一日、周窗口与失败请求处理口径。", RED),
        ("策略要求", "默认限额优先，Key override 覆盖；超限动作包括 pause、降级和通知。", BLUE),
        ("可靠性边界", "上游状态不可确认时保持现状，状态变化保留 reason 与 expires_at。", RED),
    ], "用量事实 → 成本 / Token 计算 → 策略匹配 → 治理动作 → 状态同步", "额度策略、Key 级覆盖、暂停 / 降级 / 告警、CPA 状态同步与异常保护")
    populate_quota_solution(new_inner_slide(prs, "额度治理｜实现路径", "将额度规则、Key 覆盖、超限动作和上游状态纳入同一治理链路。"))
    populate_quota_code(new_inner_slide(prs, "额度治理｜实现结构与关键逻辑", "规则匹配与上游状态分别建模，避免故障时产生不可逆的猜测性动作。"))
    populate_quota_accept(new_inner_slide(prs, "额度治理｜验收结果"))

    requirement_page(new_inner_slide(prs, "用量分析｜建设要求", "在长时间范围和高并发访问下，同时满足查询性能、数据准确与结果可解释。"), 3, "用量分析｜建设要求", "在长时间范围和高并发访问下，同时满足查询性能、数据准确与结果可解释。", [
        ("性能要求", "避免每次请求整批拉取与内存汇总；汇总、趋势和维度排名按 Tab 按需返回。", RED),
        ("数据要求", "usage_events 保留原始事实；Rollup 作为可删除、可重建的读取加速层。", BLUE),
        ("解释要求", "Rollup 未覆盖或筛选不支持时回退 raw，返回 source / complete，禁止静默少统计。", RED),
    ], "原始事实 → 聚合加速 → 按需查询 → 明细分页 → 来源与完整性说明", "usage_events、Hourly / Daily Rollup、维度聚合、Analytics include 与 Keyset 分页")
    populate_analytics_solution(new_inner_slide(prs, "用量分析｜实现路径", "以可重建数据链路支撑高并发查询，同时保持来源与完整性可解释。"))
    populate_analytics_code(new_inner_slide(prs, "用量分析｜实现结构与关键逻辑", "聚合进度只有在派生数据成功写入后推进，明细查询使用稳定 keyset。"))
    populate_analytics_accept(new_inner_slide(prs, "用量分析｜验收结果"))

    requirement_page(new_inner_slide(prs, "插件审计｜建设要求", "在扩展审计能力的同时，确保来源、方法、路径、凭证和原始数据处于受控边界。"), 4, "插件审计｜建设要求", "在扩展审计能力的同时，确保来源、方法、路径、凭证和原始数据处于受控边界。", [
        ("业务需求", "企业审计能力以插件形态独立演进，不改变核心管理面板的稳定性。", RED),
        ("安全边界", "校验 iframe origin、HTTP method、插件路径和资源 allowlist，拒绝越权访问。", BLUE),
        ("数据要求", "分析明细只提供掩码、Hash 和非秘密 metadata，不暴露 raw_json 与明文 Key。", RED),
    ], "审计需求 → 插件契约 → 来源与路径校验 → 数据脱敏 → 独立交付", "Enterprise Access Audit、插件资源管理、iframe 通信和 Analytics 安全展示")
    populate_security_solution(new_inner_slide(prs, "插件审计｜实现路径", "在扩展审计能力的同时，保持核心管理面板的稳定性与独立可用。"))
    populate_security_code(new_inner_slide(prs, "插件审计｜实现结构与关键逻辑", "通过多重边界校验将插件能力限制在可验证范围内。"))
    populate_security_accept(new_inner_slide(prs, "插件审计｜验收结果"))

    populate_summary(new_inner_slide(prs, "总结"))

    # Move the original template closing slide to the end.
    sld_id_lst = prs.slides._sldIdLst
    sld_id_lst.remove(thanks_sld_id)
    sld_id_lst.append(thanks_sld_id)

    if OUTPUT.exists(): OUTPUT.unlink()
    prs.save(str(OUTPUT))
    shutil.copy2(OUTPUT, CURRENT)
    shutil.copy2(OUTPUT, NAMED)
    print(f"created {OUTPUT}")
    print(f"updated {CURRENT}")
    print(f"updated {NAMED}")
    print(f"slides {len(prs.slides)}")


if __name__ == "__main__":
    build()
